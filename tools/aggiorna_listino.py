#!/usr/bin/env python3
"""Aggiorna il listino dentro index.html a partire dal file Excel Everdeck.

Uso:  python3 tools/aggiorna_listino.py listino.xlsx [rampe.xlsx]
Legge il foglio "Tabelle1" del listino (e, se c'e', un file rampe con un foglio per altezza)
e riscrive il blocco <script id="catalog"> di index.html.
Senza file rampe tiene i kit rampe e i loro articoli gia' presenti in index.html.
Le righe con il codice evidenziato in azzurro sono articoli integrati da fatture di acquisto:
nel catalogo hanno il campo "fattura" (nota di provenienza) e l'app li segnala nell'offerta.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

FOGLIO = "Tabelle1"
# Correzioni a errori evidenti del file (codice normalizzato -> campi da sovrascrivere)
CORREZIONI = {
    "VS75PB-1X1": {"pack": [1005, 1005, 100]},  # nel file l'imballo 1x1 e' 2005x1005x100
    # Spider deck: spessore imballo 60 mm (confermato dall'utente; nel file e' 600)
    "SPHA-2X1": {"pack": [2050, 1050, 60]}, "SPHB-2X1": {"pack": [2050, 1050, 60]},
    "SPHA-1X1": {"pack": [1050, 1050, 60]}, "SPHB-1X1": {"pack": [1050, 1050, 60]},
    # misure dai disegni quotati Everdeck: nel file sono piu' piccole del pezzo
    "VPLI": {"pack": [228, 76, 50]},       # limitatore: 228 x 76 x 50 mm (nel file 210 x 40 x 54)
    "VDT": {"pack": [2075, 608, 1188]},    # carrello verticale da porta: 2075 x 608 x 1188 mm (nel file 2060 x 600 x 1020)
}
# Articoli del file Excel da non mettere nell'app (codice normalizzato)
ESCLUSI = {"VRP20", "VRP40", "VRP60"}   # rampe da fattura: doppioni dei kit rampe, tolte su richiesta dell'utente
IMG_DIR = Path(__file__).resolve().parent.parent / "img"
AZZURRO = "BDD7EE"  # riempimento delle righe integrate da fatture di acquisto
DIM = re.compile(r"^\s*(\d+)\s*[x×]\s*(\d+)\s*[x×]\s*(\d+)\s*$", re.I)


def norm(code):
    return re.sub(r"\s+", "", code).replace("×", "x").upper()


def foto_key(name):
    """Nome foto normalizzato per l'abbinamento: senza estensione (anche 'VS750H-1x1jpg' senza punto), come norm()."""
    stem = re.sub(r"\.?(jpe?g|png|webp)\s*$", "", str(name).strip(), flags=re.I)
    return norm(stem)


def img_file(key):
    """Nome del file in img/ per una chiave foto: solo lettere, cifre, punto, trattino (la virgola diventa _)."""
    return re.sub(r"[^A-Z0-9.-]+", "_", key) + ".jpg"


def assegna_immagini(items):
    """Mette in ogni articolo il campo "img" se in img/ c'e' la sua foto (nome da colonna A, altrimenti codice)."""
    n = 0
    for i in items:
        i.pop("img", None)
        for key in (i.get("foto"), norm(i["code"])):
            if key and (IMG_DIR / img_file(key)).exists():
                i["img"] = img_file(key)
                n += 1
                break
    return n


def rampe_immagini():
    """Altezze (cm) delle rampe con immagine in img/: RAMPA-<h>.jpg (rampa) e RAMPA-<h>-PALCO.jpg (rampa montata al palco)."""
    return sorted({int(m.group(1)) for f in IMG_DIR.glob("RAMPA-*.jpg") if (m := re.match(r"RAMPA-(\d+)", f.name))})


def dims(v):
    m = DIM.match(v) if isinstance(v, str) else None
    return [int(m.group(i)) for i in (1, 2, 3)] if m else None


def num(v):
    return float(v) if isinstance(v, (int, float)) else None


def parse(xlsx):
    ws = openpyxl.load_workbook(xlsx, data_only=True)[FOGLIO]
    items, seen, sezione = [], set(), ""
    for r in range(1, ws.max_row + 1):
        code, name, price, pkgprice, kg, pdim, p1, p2 = (ws.cell(r, c).value for c in (3, 4, 5, 6, 7, 9, 10, 11))
        foto = ws.cell(r, 1).value
        fill = ws.cell(r, 3).fill
        azzurro = bool(fill.fill_type) and str(fill.fgColor.rgb or "").upper().endswith(AZZURRO)
        name = str(name).strip() if name else ""
        code = str(code).strip() if code else ""
        if code == "Product Code":
            continue
        if not code and name and price is None and kg is None:
            sezione = name
            continue
        if not name or num(price) is None:
            continue
        if code.upper().startswith(("PSP", "VSP")):  # teli: il codice 1000x2050 pleated e' scritto 1000x1050 nel file
            mm = re.search(r"(\d+x\d+)\s*$", name)
            if mm:
                code = re.sub(r"\d+x\d+\s*$", mm.group(1), code)
        if not code:  # righe senza codice nel file (profilo 5 m, tappo)
            code = "PROFILO-5M" if "profile" in name.lower() else "TAPPO-PLASTICA"
        key = norm(code)
        if key in ESCLUSI or key in seen:
            continue
        seen.add(key)
        item = {
            "code": code,
            "name": re.sub(r"\s+", " ", name),
            "section": sezione,
            "price": round(num(price), 4),
            "kg": num(kg),
            "pack": dims(p1) or dims(p2),
        }
        # Articoli venduti in confezione (gambe x4, connettori x3): "Package price" = n x prezzo unitario.
        # Peso e imballo del file sono della confezione, il prezzo unitario e' per pezzo.
        if num(pkgprice) and num(price) and dims(p2) and not dims(p1):
            n = num(pkgprice) / num(price)
            if n > 1.05 and abs(n - round(n)) < 0.05:
                item["pkg"] = int(round(n))
        if foto and str(foto).strip() and str(foto).strip() != "Product Photo Name":
            item["foto"] = foto_key(foto)
        if azzurro:
            # colonna R: provenienza (codice, fatture), colonna U: come sono stati ricavati peso e misure
            nota = " · ".join(re.sub(r"\s+", " ", str(v)).strip() for v in (ws.cell(r, 18).value, ws.cell(r, 21).value) if v)
            item["fattura"] = nota or "articolo integrato da fatture di acquisto"
        item.update(CORREZIONI.get(key, {}))
        items.append(item)
    return items


def parse_rampe(xlsx, gia):
    """Ogni foglio e' un kit rampa: righe con codice, descrizione, quantita', prezzo."""
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    articoli, kit = [], []
    for ws in wb:
        numeri = [int(n) for n in re.findall(r"\d+", ws.title)]
        if not numeri:
            continue
        parti = []
        for r in range(1, ws.max_row + 1):
            code, desc, qty, price = (ws.cell(r, c).value for c in (2, 3, 4, 5))
            if not code or code == "Part number" or num(price) is None or num(qty) is None:
                continue
            code = str(code).strip()
            parti.append({"code": code, "qty": int(num(qty))})
            key = norm(code)
            if key not in gia:
                gia.add(key)
                articoli.append({"code": code, "name": re.sub(r"\s+", " ", str(desc).strip()), "section": "RAMPE",
                                 "price": round(num(price), 4), "kg": None, "pack": None})
        kit.append({"name": ws.title.strip(), "lo": numeri[0], "hi": numeri[-1], "parts": parti})
    return articoli, kit


def main():
    if len(sys.argv) not in (2, 3):
        sys.exit(__doc__)
    xlsx = Path(sys.argv[1])
    items = parse(xlsx)
    html_path = Path(__file__).resolve().parent.parent / "index.html"
    html = html_path.read_text(encoding="utf-8")
    pat = re.compile(r'(<script type="application/json" id="catalog">)(.*?)(</script>)', re.S)
    vecchio = pat.search(html)
    if not vecchio:
        sys.exit("Blocco catalog non trovato in index.html")
    gia = {norm(i["code"]) for i in items}
    if len(sys.argv) == 3:
        extra, kit = parse_rampe(sys.argv[2], gia)
    else:  # nessun file rampe: tengo kit e articoli rampe gia' in index.html
        old = json.loads(vecchio.group(2).replace("<\\/", "</"))
        kit = old.get("rampKits", [])
        extra = [i for i in old["items"] if i.get("section") == "RAMPE" and norm(i["code"]) not in gia]
        if kit:
            print(f"Kit rampe tenuti da index.html: {len(kit)} kit, {len(extra)} articoli (prezzi non inclusi in prezzi.json).")
    items += extra
    n_img = assegna_immagini(items)
    m = re.search(r"\d{4}\.\d{2}", xlsx.stem)
    label = f"Everdeck STAGEDECKS {m.group(0)}" if m else xlsx.stem
    # i prezzi NON vanno nel repository: restano in prezzi.json (ignorato da git)
    prezzi = {norm(i["code"]): i["price"] for i in items if i.get("price") is not None}
    # prezzi che non sono nel listino Excel (es. pezzi speciali dei palchi rotondi): prezzi_extra.json, anche lui fuori da git
    extra = Path(__file__).resolve().parent.parent / "prezzi_extra.json"
    if extra.exists():
        prezzi.update({norm(k): v for k, v in json.loads(extra.read_text(encoding="utf-8")).items() if not k.startswith("_")})
    for i in items:
        i.pop("price", None)
    (Path(__file__).resolve().parent.parent / "prezzi.json").write_text(
        json.dumps({"listino": label, "prezzi": prezzi}, ensure_ascii=False, indent=0), encoding="utf-8")
    payload = json.dumps({"listino": label, "items": items, "rampKits": kit, "rampImgs": rampe_immagini()}, ensure_ascii=False, separators=(",", ":"))
    payload = payload.replace("</", "<\\/")
    html_path.write_text(pat.sub(lambda mo: mo.group(1) + payload + mo.group(3), html, count=1), encoding="utf-8")
    print(f"{len(items)} articoli scritti in index.html ({label}), senza prezzi.")
    print(f"Immagini trovate in img/: {n_img} articoli.")
    fatt = [i["code"] for i in items if i.get("fattura")]
    if fatt:
        print(f"{len(fatt)} articoli integrati da fatture (righe azzurre): {', '.join(fatt)}")
    print("Prezzi salvati in prezzi.json: NON caricarlo su GitHub, caricalo nell'app (scheda Listino).")


if __name__ == "__main__":
    main()
